"""任务编排：把「翻译轨」与「课程知识轨」串起来。

    .pptx  ─┬─ 文本段落翻译（旧版行为：递归组合图形、表格单元格、并发、写回首个 run）
            └─ 图片内文字翻译（OCR → 擦除原文 → 重绘译文 / 叠加可编辑文本框）
    .pdf   ─┬─ 离线解析（PyMuPDF）→ 翻译文本与表格 → 重建可编辑 PPTX（背景位图保版式）
            └─ 再对页面背景里的图内文字做一遍 OCR 翻译（叠加可编辑文本框）

两条轨都产出 <原名>_<目标语言>.pptx，命名习惯与旧版一致（如 Week8_TES_Chinese.pptx）。
"""

from __future__ import annotations

import hashlib
import io
import os
from concurrent.futures import ThreadPoolExecutor, as_completed
from dataclasses import dataclass, field

from deepseek_client import DeepSeekClient, should_translate
from pptx_images import (
    add_overlay_textbox,
    collect_pictures,
    map_pixel_box,
    picture_bytes,
    replace_picture_image,
)
from translator_config import TranslatorConfig, output_path_for


@dataclass(slots=True)
class JobResult:
    output_path: str = ""
    mode: str = ""
    slides: int = 0
    text_units: int = 0
    translated_units: int = 0
    tables: int = 0
    failures: int = 0
    images_total: int = 0
    images_translated: int = 0
    image_regions: int = 0
    notes: list[str] = field(default_factory=list)
    knowledge_dir: str = ""
    report_files: list[str] = field(default_factory=list)
    outline_points: int = 0
    transcript_segments: int = 0
    exam_signals: int = 0


class UnsupportedInput(RuntimeError):
    pass


def unique_output_path(path: str) -> str:
    """绝不覆盖已有文件：同名时自动追加 _2 / _3 后缀。

    这条规则由来于一次真实事故：开发期一次端到端测试把输出写进了用户目录，
    与用户已有的《Lecture03_..._Chinese.pptx》同名，直接覆盖掉了别人的成品。
    因此本工具默认「只新增、不覆盖」；确实需要覆盖时必须显式传 output_path。
    """
    if not os.path.exists(path):
        return path
    stem, extension = os.path.splitext(path)
    index = 2
    while os.path.exists(f"{stem}_{index}{extension}"):
        index += 1
    return f"{stem}_{index}{extension}"


def make_client(config: TranslatorConfig, log=None) -> DeepSeekClient:
    return DeepSeekClient(
        config.api_key,
        config.effective_model(),
        api_base=config.api_base,
        timeout=config.request_timeout,
        log=log,
    )


def _translate_texts(texts, config, *, log=None, progress=None, context="", style="auto"):
    """并发翻译一批文本；失败项保留原文，保证不丢内容。"""
    if config.only_convert or not texts:
        return list(texts)
    client = make_client(config, log=log)
    results = list(texts)
    with ThreadPoolExecutor(max_workers=max(1, int(config.max_workers))) as pool:
        futures = {
            pool.submit(
                client.translate_batch,
                chunk,
                src_lang=config.source_language_name(),
                tgt_lang=config.target_language_name(),
                context=context,
                style=style,
            ): index
            for index, chunk in enumerate(_batched(texts))
        }
        done = 0
        for future in as_completed(futures):
            index = futures[future]
            done += 1
            try:
                translated = future.result()
                offset = index * BATCH_SIZE
                for position, value in enumerate(translated):
                    if position + offset < len(results) and value:
                        results[position + offset] = value
            except Exception as exc:  # noqa: BLE001
                if log:
                    log(f"批量翻译失败（保留原文）：{exc}")
            if progress:
                progress(done, len(futures))
    return results


BATCH_SIZE = 12


def _batched(items, size: int = BATCH_SIZE):
    for start in range(0, len(items), size):
        yield list(items[start : start + size])


# ---------------------------------------------------------------------- 图片轨
def _image_key(blob: bytes) -> str:
    return hashlib.sha1(blob).hexdigest()


def translate_images_in_pptx(prs, config: TranslatorConfig, client, *, log, progress=None) -> tuple[int, int, int]:
    """处理整个演示文稿里的图片文字。

    分三段跑，避免把网络往返串到 CPU 流水线上：
      A 串行 OCR（CPU 密集）→ B 并发翻译（网络等待）→ C 串行重绘与写回。

    返回 ``(图片总数, 处理过的图片数, 重绘文字块数)``。
    """
    from image_translator import prepare_image, render_plan

    refs = collect_pictures(prs)
    if not refs:
        log("没有发现图片，跳过图片文字翻译")
        return 0, 0, 0

    # ---- A：OCR（按内容去重，同一张图只识别一次）----
    plans: dict[str, object] = {}
    shape_map: dict[str, list] = {}
    scanned = set()
    for image_index, ref in enumerate(refs, 1):
        blob, _ext = picture_bytes(ref.shape)
        if not blob:
            continue
        key = _image_key(blob)
        shape_map.setdefault(key, []).append(ref)
        if key in scanned:
            continue
        scanned.add(key)
        log(f"图片识别 {image_index}/{len(refs)}（含小字分块扫描）…")
        try:
            plan = prepare_image(
                io.BytesIO(blob),
                min_confidence=config.ocr_min_confidence,
                min_chars=config.ocr_min_chars,
                max_side=config.ocr_max_side,
                should_translate=should_translate,
            )
        except Exception as exc:  # noqa: BLE001
            log(f"图片 OCR 失败（跳过）：{exc}")
            plan = None
        if plan is not None:
            plans[key] = plan
            log(f"第 {ref.slide_index + 1} 页：识别到 {len(plan.lines)} 处待翻译图中文字")
    if not plans:
        log(f"扫描了 {len(shape_map)} 张图片，没有发现需要翻译的文字")
        return len(shape_map), 0, 0

    total_regions = sum(len(plan.lines) for plan in plans.values())
    log(f"图片文字翻译：{len(plans)} 张图命中文字，共 {total_regions} 个文字块，开始翻译…")

    # ---- B：并发翻译 ----
    keys = list(plans.keys())
    translations: dict[str, list[str]] = {}
    workers = max(1, min(int(config.max_workers), 8))

    def work(key: str):
        plan = plans[key]
        return client.translate_batch(
            plan.texts(),
            src_lang=config.source_language_name(),
            tgt_lang=config.target_language_name(),
            context="学术课件插图（示意图/图表/坐标轴/图注）里的短标签",
            style="label",
        )

    with ThreadPoolExecutor(max_workers=workers) as pool:
        futures = {pool.submit(work, key): key for key in keys}
        done = 0
        for future in as_completed(futures):
            key = futures[future]
            done += 1
            try:
                translations[key] = future.result()
            except Exception as exc:  # noqa: BLE001
                log(f"图片文字翻译失败（保留原图）：{exc}")
                translations[key] = []
            if progress:
                progress(done, len(futures))

    # ---- C：重绘 / 叠加写回 ----
    processed = 0
    region_count = 0
    for key, plan in plans.items():
        values = translations.get(key) or []
        if len(values) != len(plan.lines):
            continue
        try:
            image, stats = render_plan(
                plan, values, mode=config.image_mode, font_path=None, log=log
            )
        except Exception as exc:  # noqa: BLE001
            log(f"图片重绘失败（保留原图）：{exc}")
            continue
        if image is None or not stats.regions_detail:
            continue
        log(stats.note)
        buffer = io.BytesIO()
        image.save(buffer, format="PNG")
        payload = buffer.getvalue()
        source_size = plan.pil.size
        for ref in shape_map.get(key, []):
            original_shapes = set(shape.shape_id for shape in ref.slide.shapes)
            if config.image_mode != "inpaint":
                # 有旋转/裁剪无法映射时必须保留原图，不能先擦除再静默跳过文本框。
                if not _add_overlay_textboxes(ref, source_size, stats.regions_detail, log):
                    continue
            if not replace_picture_image(ref.shape, payload):
                for shape in list(ref.slide.shapes):
                    if shape.shape_id not in original_shapes:
                        shape._element.getparent().remove(shape._element)
                log("图片写回失败（保留原图）")
                continue
            processed += 1
            region_count += len(stats.regions_detail)
    log(f"图片文字翻译完成：{processed} 张图 / {region_count} 处文字")
    return len(shape_map), processed, region_count


def _add_overlay_textboxes(ref, image_size, regions, log) -> bool:
    boxes = [map_pixel_box(ref.shape, image_size, region.box) for region in regions]
    if any(box is None for box in boxes):
        log("图片旋转或坐标无法映射，保留原图；可使用重绘图片模式")
        return False
    added = []
    for region, box in zip(regions, boxes):
        left, top, width, height = box
        size_pt = max(6.0, min(40.0, height / 12700.0 * 0.72))
        try:
            added.append(add_overlay_textbox(
                ref.slide,
                left=left,
                top=top,
                width=width,
                height=height,
                text=region.translated_text,
                font_size_pt=size_pt,
                color=region.color,
                vertical=region.vertical,
            ))
        except Exception as exc:  # noqa: BLE001
            for shape in added:
                shape._element.getparent().remove(shape._element)
            log(f"叠加文本框失败（保留原图）：{exc}")
            return False
    return True


def make_deck_image_translator(config: TranslatorConfig, client, log):
    """给课程知识库用的插图翻译：把 02_slides/ 里的关键图就地翻成中文。"""
    from image_translator import translate_image

    def translate(deck_path: str, deck) -> None:
        targets = [(slide, path) for slide in deck.slides for path in slide.images]
        if not targets:
            return
        if config.only_convert:
            return
        log(f"课件插图翻译：{len(targets)} 张关键图")
        done = 0
        for slide, path in targets:
            try:
                image, stats = translate_image(
                    path,
                    lambda texts, context="", style="label": client.translate_batch(
                        texts,
                        src_lang=config.source_language_name(),
                        tgt_lang=config.target_language_name(),
                        context="学术课件插图里的短标签",
                        style="label",
                    ),
                    mode="inpaint",
                    min_confidence=config.ocr_min_confidence,
                    min_chars=config.ocr_min_chars,
                    max_side=config.ocr_max_side,
                    should_translate=should_translate,
                    log=log,
                )
            except Exception as exc:  # noqa: BLE001
                log(f"插图翻译失败（{os.path.basename(path)}）：{exc}")
                continue
            if image is None:
                continue
            extension = os.path.splitext(path)[1] or ".png"
            try:
                if extension.lower() in (".jpg", ".jpeg"):
                    image.convert("RGB").save(path, quality=92)
                else:
                    image.save(path)
                done += 1
            except Exception as exc:  # noqa: BLE001
                log(f"插图写回失败（{os.path.basename(path)}）：{exc}")
        log(f"课件插图翻译完成：{done}/{len(targets)} 张已改写为中文")

    return translate


# ---------------------------------------------------------------------- 翻译轨
def translate_pptx(src_path: str, dst_path: str, config: TranslatorConfig, *, log, progress=None) -> JobResult:
    from pptx import Presentation

    from pptx_translator import apply_translation, collect_text_units

    result = JobResult(output_path=dst_path, mode="PPTX 文本 + 图片")
    prs = Presentation(src_path)
    result.slides = len(prs.slides)

    units = collect_text_units(prs)
    result.text_units = len(units)
    log(f"共发现 {len(units)} 个有效待翻译文本段落")

    if not config.only_convert and units:
        client = make_client(config, log=log)

        def translate(paragraph_unit):
            return client.translate(
                paragraph_unit.text,
                src_lang=config.source_language_name(),
                tgt_lang=config.target_language_name(),
            )

        results: dict[int, str] = {}
        with ThreadPoolExecutor(max_workers=max(1, int(config.max_workers))) as pool:
            futures = {pool.submit(translate, unit): index for index, unit in enumerate(units)}
            done = 0
            for future in as_completed(futures):
                index = futures[future]
                done += 1
                try:
                    results[index] = future.result()
                    result.translated_units += 1
                except Exception as exc:  # noqa: BLE001
                    result.failures += 1
                    log(f"翻译失败（保留原文）[{index}]：{exc}")
                if progress:
                    progress(done * 0.7, len(units))
        for index, unit in enumerate(units):
            translated = results.get(index)
            if translated:
                apply_translation(unit.paragraph, translated)

    if config.translate_images:
        client = make_client(config, log=log)
        total, processed, regions = translate_images_in_pptx(
            prs, config, client, log=log, progress=None
        )
        result.images_total = total
        result.images_translated = processed
        result.image_regions = regions

    prs.save(dst_path)
    if progress:
        progress(1.0, 1.0)
    return result


def translate_pdf(src_path: str, dst_path: str, config: TranslatorConfig, *, log, progress=None) -> JobResult:
    from pdf_builder import build_pptx_from_pdf
    from pdf_parser import parse_pdf

    result = JobResult(output_path=dst_path, mode="PDF → PPT 后翻译")
    log("正在解析 PDF（离线，不依赖任何云 API）…")
    parsed = parse_pdf(src_path, detect_tables=config.pdf_detect_tables)
    log(
        f"解析完成：{len(parsed.pages)} 页 / "
        f"{parsed.text_element_count} 个文本块 / {parsed.table_count} 个表格"
    )
    if not parsed.pages:
        raise UnsupportedInput("PDF 没有可处理的页面")
    if parsed.text_element_count == 0 and parsed.table_count == 0 and not (
        config.only_convert or (config.translate_images and config.pdf_ocr_images and config.pdf_rasterize_background)
    ):
        raise UnsupportedInput(
            "这个 PDF 里没有可提取的文本（可能是扫描图片版）。"
            "请启用图片翻译、PDF 图片 OCR 和页面背景选项。"
        )
    result.slides = len(parsed.pages)
    result.tables = parsed.table_count

    texts: list[str] = []
    text_refs: list[tuple[int, int]] = []
    cell_texts: list[str] = []
    cell_refs: list[tuple[int, int, int, int]] = []

    for page_index, page in enumerate(parsed.pages):
        for element_index, element in enumerate(page.elements):
            if element.kind == "table" and element.table:
                for row_index, row in enumerate(element.table):
                    for col_index, value in enumerate(row):
                        if should_translate(value):
                            cell_texts.append(value.strip())
                            cell_refs.append((page_index, element_index, row_index, col_index))
                continue
            if should_translate(element.text):
                texts.append(element.text.strip())
                text_refs.append((page_index, element_index))

    result.text_units = len(texts) + len(cell_texts)
    log(f"待翻译：{len(texts)} 个文本块 + {len(cell_texts)} 个表格单元格")

    if not config.only_convert and result.text_units:
        translated_texts = _translate_texts(
            texts, config, log=log, context="学术讲义 PDF 页面文本", style="prose"
        )
        translated_cells = _translate_texts(
            cell_texts, config, log=log, context="学术讲义 PDF 表格单元格", style="label"
        )
        for (page_index, element_index), value in zip(text_refs, translated_texts):
            parsed.pages[page_index].elements[element_index].translated = value
        for (page_index, element_index, row_index, col_index), value in zip(cell_refs, translated_cells):
            element = parsed.pages[page_index].elements[element_index]
            if element.table_translated is None:
                element.table_translated = [list(row) for row in (element.table or [])]
            try:
                element.table_translated[row_index][col_index] = value
            except IndexError:
                pass
        result.translated_units = result.text_units

    stats = build_pptx_from_pdf(
        parsed,
        dst_path,
        target_lang=config.tgt_lang,
        dpi=config.pdf_dpi,
        keep_page_size=config.pdf_keep_page_size,
        rasterize_background=config.pdf_rasterize_background,
        progress=lambda done, total: progress(done / max(1, total) * 0.8, total) if progress else None,
    )
    result.notes.append(f"背景页 {stats.background_pages} / 文本框 {stats.text_boxes} / 表格 {stats.tables}")

    if config.translate_images and config.pdf_ocr_images and not config.only_convert:
        log("对 PDF 页面里的图内文字再做一遍 OCR 翻译…")
        from pptx import Presentation

        prs = Presentation(dst_path)
        client = make_client(config, log=log)
        total, processed, regions = translate_images_in_pptx(prs, config, client, log=log)
        result.images_total = total
        result.images_translated = processed
        result.image_regions = regions
        if processed:
            prs.save(dst_path)
            result.notes.append(f"页面图内文字：{regions} 处已按原位置处理")

    if progress:
        progress(1.0, 1.0)
    return result


def run_translate_job(
    src_path: str,
    config: TranslatorConfig,
    *,
    log=None,
    progress=None,
    output_path: str | None = None,
) -> JobResult:
    log = log or (lambda _m: None)
    progress = progress or (lambda *_a: None)
    extension = os.path.splitext(src_path)[1].lower()
    if extension not in (".pptx", ".pdf"):
        raise UnsupportedInput(f"不支持的文件类型：{extension}（只支持 .pptx 与 .pdf）")

    if output_path:
        destination = output_path
    else:
        tag = "Converted" if config.only_convert else config.target_language_name()
        base = output_path_for(src_path, config, suffix=tag)
        if config.output_dir:
            os.makedirs(config.output_dir, exist_ok=True)
            base = os.path.join(config.output_dir, os.path.basename(base))
        destination = unique_output_path(base)
    os.makedirs(os.path.dirname(os.path.abspath(destination)) or ".", exist_ok=True)

    if extension == ".pptx":
        return translate_pptx(src_path, destination, config, log=log, progress=progress)
    return translate_pdf(src_path, destination, config, log=log, progress=progress)


# ---------------------------------------------------------------------- 课程知识轨
def default_course_dir(deck_path: str, config: TranslatorConfig) -> str:
    if config.course_dir:
        return config.course_dir
    stem = os.path.splitext(os.path.basename(deck_path))[0]
    return os.path.join(os.path.dirname(os.path.abspath(deck_path)), f"{stem}_课程知识库")


def run_course_job(
    deck_path: str,
    materials: list[str],
    config: TranslatorConfig,
    *,
    log=None,
    progress=None,
    explicit_order: bool = False,
    course_dir: str | None = None,
):
    """跑课程知识流水线；``materials`` 为空时只做课件轨（知识点总结）。"""
    import course_knowledge

    log = log or (lambda _m: None)
    client = make_client(config, log=log)
    target_dir = course_dir or default_course_dir(deck_path, config)
    translator = (
        make_deck_image_translator(config, client, log)
        if (config.translate_images and not config.only_convert)
        else None
    )
    result = course_knowledge.run_course_job(
        deck_path,
        materials,
        config,
        client=client,
        course_dir=target_dir,
        log=log,
        progress=progress,
        explicit_order=explicit_order,
        translate_images_fn=translator,
    )
    study_docs = export_study_documents(target_dir, materials=bool(materials), config=config, log=log)
    for path in study_docs:
        result.add(path)
    return result


def export_study_documents(course_dir: str, *, materials: bool, config, log=None) -> list[str]:
    """把 Markdown 产物合成为给学生看的 Word 文档。"""
    from markdown_doc import render_markdown_to_docx

    log = log or (lambda _m: None)
    notes_dir = os.path.join(course_dir, "07_notes")

    def read(name: str) -> str:
        path = os.path.join(notes_dir, name)
        if not os.path.exists(path):
            return ""
        with open(path, "r", encoding="utf-8") as handle:
            return handle.read()

    course = config.course_name or "本讲"
    produced: list[str] = []

    summary = read("summary.md")
    notes = read("notes.md")
    if summary or notes:
        body = []
        if summary:
            body.append("# 一页速览\n\n" + summary)
        if notes:
            body.append("# 详细内容\n\n" + notes)
        out = os.path.join(course_dir, f"{course}_知识点总结.docx")
        try:
            render_markdown_to_docx(
                "\n\n---\n\n".join(body),
                out,
                title=f"{course} 知识点总结",
                subtitle="由「PPT 课件学习助手」自动生成 · 所有结论均可溯源到课件页码或录音时间戳",
                images_base_dir=os.path.join(course_dir, "02_slides"),
            )
            produced.append(out)
        except Exception as exc:  # noqa: BLE001
            log(f"知识点总结 Word 生成失败：{exc}")

    if materials:
        sections: list[str] = []
        for heading, name in (
            ("考试重点", "exam-focus.md"),
            ("讲课流程", "lecture-flow.md"),
            ("详细内容报告", "notes.md"),
            ("一页速览", "summary.md"),
        ):
            content = read(name)
            if content:
                sections.append(f"# {heading}\n\n{content}")
        if sections:
            out = os.path.join(course_dir, f"{course}_课堂复习报告.docx")
            try:
                render_markdown_to_docx(
                    "\n\n---\n\n".join(sections),
                    out,
                    title=f"{course} 课堂复习报告",
                    subtitle="考试重点已按老师原话标注来源 · 详见各节溯源标记",
                    images_base_dir=os.path.join(course_dir, "02_slides"),
                )
                produced.append(out)
            except Exception as exc:  # noqa: BLE001
                log(f"课堂复习报告 Word 生成失败：{exc}")
    return produced
