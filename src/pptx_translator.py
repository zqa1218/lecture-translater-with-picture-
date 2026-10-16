"""PPTX 文本翻译：递归收集段落 + 把译文写回。

对齐旧版「全球通翻译工具」的行为：

  * 递归遍历所有 shape，包括组合图形（Group Shape）内的子 shape；
  * 文本走 text_frame 的段落，表格走 table 的每个单元格；
  * 讲者备注（Notes）一并翻译，但**只在该页本来就有备注页时才碰**，
    避免给没有备注的幻灯片凭空造出空的备注页；
  * 译文写入段落的第一个 run，其余 run 清空 —— 保留原字体与段落版式。

并发与写出由 ``pipeline.translate_pptx`` 负责，本模块只做「收集」与「写回」。
"""

from __future__ import annotations

from dataclasses import dataclass

from pptx.enum.shapes import MSO_SHAPE_TYPE

from deepseek_client import should_translate


@dataclass(slots=True)
class TextUnit:
    paragraph: object
    text: str


def _iter_paragraphs(text_frame):
    for paragraph in text_frame.paragraphs:
        yield paragraph


def iter_text_units(shapes):
    """递归产出所有可翻译段落（含组合图形与表格单元格）。"""
    for shape in shapes:
        shape_type = getattr(shape, "shape_type", None)
        if shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from iter_text_units(shape.shapes)
            continue

        if getattr(shape, "has_table", False):
            try:
                table = shape.table
            except Exception:  # noqa: BLE001
                continue
            for row in table.rows:
                for cell in row.cells:
                    yield from _iter_paragraphs(cell.text_frame)
            continue

        if getattr(shape, "has_text_frame", False):
            try:
                frame = shape.text_frame
            except Exception:  # noqa: BLE001
                continue
            yield from _iter_paragraphs(frame)


def iter_slide_paragraphs(slide):
    """一页的全部段落：正文（含组合图形/表格）+ 原有备注。"""
    yield from iter_text_units(slide.shapes)
    if getattr(slide, "has_notes_slide", False):
        yield from iter_text_units(slide.notes_slide.shapes)


def collect_text_units(prs) -> list[TextUnit]:
    units: list[TextUnit] = []
    for slide in prs.slides:
        for paragraph in iter_slide_paragraphs(slide):
            text = paragraph.text or ""
            if should_translate(text):
                units.append(TextUnit(paragraph=paragraph, text=text.strip()))
    return units


def apply_translation(paragraph, translated: str) -> None:
    """写入第一个 run，清空其余 run（保留原格式）。"""
    runs = list(paragraph.runs)
    if not runs:
        paragraph.add_run().text = translated
        return
    runs[0].text = translated
    for extra in runs[1:]:
        extra.text = ""