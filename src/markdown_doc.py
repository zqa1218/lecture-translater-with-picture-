"""Markdown → Word(.docx) 渲染。

为什么自己写而不装 pandoc：pandoc 是外部程序，没法塞进单文件 exe；
而 python-docx 只需要几百 KB。这里覆盖课程产物实际用到的语法：
标题、段落、粗体/斜体/行内代码、项目符号、编号列表、表格、引用、代码块、水平线、图片。
"""

from __future__ import annotations

import os
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

BODY_FONT = "宋体"
HEADING_FONT = "微软雅黑"
MONO_FONT = "Consolas"

INLINE_RE = re.compile(r"(\*\*.+?\*\*|__.+?__|`[^`]+`|\*[^*\n]+?\*)")
HEADING_RE = re.compile(r"^(#{1,6})\s+(.*)$")
BULLET_RE = re.compile(r"^\s*[-*+]\s+(.*)$")
ORDERED_RE = re.compile(r"^\s*(\d{1,3})[.)、]\s+(.*)$")
IMAGE_RE = re.compile(r"^!\[(?P<alt>[^\]]*)\]\((?P<path>[^)]+)\)\s*$")
SEPARATOR_RE = re.compile(r"^\s*\|?[\s:|-]{3,}\|?\s*$")
FENCE_RE = re.compile(r"^\s*```")
HR_RE = re.compile(r"^\s*(-{3,}|\*{3,}|_{3,})\s*$")


def _style_run(run, *, font: str = BODY_FONT, size: float | None = None, color=None) -> None:
    run.font.name = font
    if size:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor(*color)
    rpr = run._element.get_or_add_rPr()  # noqa: SLF001
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:eastAsia"), font)
    rfonts.set(qn("w:ascii"), font)
    rfonts.set(qn("w:hAnsi"), font)


def _add_inline(paragraph, text: str) -> None:
    for token in INLINE_RE.split(text):
        if not token:
            continue
        if token.startswith("**") and token.endswith("**") and len(token) > 4:
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("__") and token.endswith("__") and len(token) > 4:
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("`") and token.endswith("`") and len(token) > 2:
            run = paragraph.add_run(token[1:-1])
            _style_run(run, font=MONO_FONT, size=9)
            continue
        elif token.startswith("*") and token.endswith("*") and len(token) > 2:
            run = paragraph.add_run(token[1:-1])
            run.italic = True
        else:
            run = paragraph.add_run(token)
        _style_run(run)


def _configure_styles(document: Document) -> None:
    normal = document.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(10.5)
    rpr = normal.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:eastAsia"), BODY_FONT)
    for level in range(1, 5):
        try:
            style = document.styles[f"Heading {level}"]
        except KeyError:
            continue
        style.font.name = HEADING_FONT
        style.font.color.rgb = RGBColor(0x1F, 0x3B, 0x63)
        style_rpr = style.element.get_or_add_rPr()
        style_fonts = style_rpr.find(qn("w:rFonts"))
        if style_fonts is None:
            style_fonts = style_rpr.makeelement(qn("w:rFonts"), {})
            style_rpr.insert(0, style_fonts)
        style_fonts.set(qn("w:eastAsia"), HEADING_FONT)


def _table_rows(lines: list[str]) -> list[list[str]]:
    rows: list[list[str]] = []
    for line in lines:
        cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
        rows.append(cells)
    return rows


def render_markdown_to_docx(
    markdown_text: str,
    output_path: str,
    *,
    title: str | None = None,
    images_base_dir: str | None = None,
    subtitle: str | None = None,
) -> str:
    document = Document()
    _configure_styles(document)

    if title:
        heading = document.add_heading(title, level=0)
        for run in heading.runs:
            _style_run(run, font=HEADING_FONT, size=20)
    if subtitle:
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(subtitle)
        _style_run(run, size=9, color=(0x66, 0x66, 0x66))

    lines = markdown_text.replace("\r\n", "\n").split("\n")
    index = 0
    while index < len(lines):
        line = lines[index]

        if FENCE_RE.match(line):
            index += 1
            buffer: list[str] = []
            while index < len(lines) and not FENCE_RE.match(lines[index]):
                buffer.append(lines[index])
                index += 1
            index += 1
            for chunk in buffer:
                paragraph = document.add_paragraph()
                run = paragraph.add_run(chunk)
                _style_run(run, font=MONO_FONT, size=9)
                paragraph.paragraph_format.space_after = Pt(0)
            continue

        if not line.strip():
            index += 1
            continue

        if HR_RE.match(line):
            paragraph = document.add_paragraph()
            run = paragraph.add_run("─" * 30)
            _style_run(run, size=8, color=(0xAA, 0xAA, 0xAA))
            index += 1
            continue

        heading = HEADING_RE.match(line)
        if heading:
            level = min(4, len(heading.group(1)))
            text = heading.group(2).strip()
            document.add_heading("", level=level)
            paragraph = document.paragraphs[-1]
            _add_inline(paragraph, text)
            for run in paragraph.runs:
                _style_run(run, font=HEADING_FONT, size={1: 16, 2: 14, 3: 12, 4: 11}[level])
                run.bold = True
            index += 1
            continue

        image = IMAGE_RE.match(line)
        if image:
            path = image.group("path").strip()
            if images_base_dir and not os.path.isabs(path):
                path = os.path.join(images_base_dir, path)
            if os.path.exists(path):
                try:
                    document.add_picture(path, width=Inches(5.8))
                except Exception:  # noqa: BLE001
                    pass
            else:
                run = document.add_paragraph().add_run(f"[图片缺失] {image.group('path')}")
                _style_run(run, size=9, color=(0xAA, 0xAA, 0xAA))
            index += 1
            continue

        if line.lstrip().startswith("|") and "|" in line.strip()[1:]:
            block: list[str] = []
            while index < len(lines) and lines[index].lstrip().startswith("|"):
                if not SEPARATOR_RE.match(lines[index]):
                    block.append(lines[index])
                index += 1
            rows = _table_rows(block)
            if rows:
                columns = max(len(row) for row in rows)
                table = document.add_table(rows=len(rows), cols=columns)
                table.style = "Table Grid"
                for row_index, row in enumerate(rows):
                    for col_index in range(columns):
                        cell = table.cell(row_index, col_index)
                        value = row[col_index] if col_index < len(row) else ""
                        cell.text = ""
                        paragraph = cell.paragraphs[0]
                        _add_inline(paragraph, value)
                        for run in paragraph.runs:
                            _style_run(run, size=9)
                            if row_index == 0:
                                run.bold = True
            continue

        if line.lstrip().startswith(">"):
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.3)
            _add_inline(paragraph, line.lstrip()[1:].strip())
            for run in paragraph.runs:
                run.italic = True
            index += 1
            continue

        bullet = BULLET_RE.match(line)
        if bullet:
            paragraph = document.add_paragraph(style="List Bullet")
            _add_inline(paragraph, bullet.group(1))
            index += 1
            continue

        ordered = ORDERED_RE.match(line)
        if ordered:
            paragraph = document.add_paragraph(style="List Number")
            _add_inline(paragraph, ordered.group(2))
            index += 1
            continue

        paragraph = document.add_paragraph()
        _add_inline(paragraph, line.strip())
        index += 1

    directory = os.path.dirname(os.path.abspath(output_path))
    if directory:
        os.makedirs(directory, exist_ok=True)
    document.save(output_path)
    return output_path
