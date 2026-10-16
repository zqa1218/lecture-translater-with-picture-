"""课程知识轨：大纲解析、考试信号扫描、分段、课件逐页提取、Word 渲染。"""

from __future__ import annotations

import io
import os

from PIL import Image

import course_knowledge as ck
from markdown_doc import render_markdown_to_docx
from transcript_io import Segment

OUTLINE = """## KP01 注意与执行控制
### KP01-01 冲突监测
- [KP01-01-01] Stroop 任务与冲突代价 ｜ 理解 ｜ PPT p.22 + 录音 38:10
  定义：不一致条件与一致条件反应时之差
- [KP01-01-02] ERN 与错误监测 ｜ 应用 ｜ PPT p.25
  定义：错误发生后 50~100 ms 的负波
### KP01-02 执行功能
- [KP01-02-01] 抑制控制 ｜ 了解 ｜ PPT p.30
  定义：压制优势反应的能力
## KP02 工作记忆
### KP02-01 模型
- [KP02-01-01] 中央执行系统 ｜ 理解 ｜ PPT p.40
  定义：工作记忆的注意控制中枢
"""


def test_parse_outline_markdown_builds_hierarchy():
    parsed = ck.parse_outline_markdown(OUTLINE)
    assert parsed["pointCount"] == 4
    units = [unit["title"] for unit in parsed["units"]]
    assert units == ["KP01 注意与执行控制", "KP02 工作记忆"]
    first = parsed["units"][0]["sections"][0]["points"][0]
    assert first["id"] == "KP01-01-01"
    assert first["level"] == "理解"
    assert first["sources"] == ["PPT p.22", "录音 38:10"]
    assert "反应时之差" in first["definition"]


def test_parse_outline_handles_empty_text():
    parsed = ck.parse_outline_markdown("")
    assert parsed == {"units": [], "pointCount": 0}


def test_scan_exam_signals_finds_keywords():
    segments = [
        Segment(text="This will be on the exam.", start=5),
        Segment(text="我们先复习一下上次的内容。", start=30),
        Segment(text="今天天气不错。", start=60),
        Segment(text="Please review chapter four before next week.", start=90),
    ]
    hits = ck.scan_exam_signals(segments)
    assert len(hits) == 3
    assert hits[0][0] == 5
    assert "exam" in hits[0][2].lower()


def test_chunk_segments_respects_char_budget():
    segments = [Segment(text="x" * 100, start=i) for i in range(100)]
    chunks = ck._chunk_segments(segments, 1000)
    assert len(chunks) > 1
    assert sum(len(c) for c in chunks) == 100


def test_chunk_slides_keeps_pages_intact():
    slides = [ck.SlideContent(index=i, texts=["y" * 100]) for i in range(1, 21)]
    chunks = ck._chunk_slides(slides, 1000)
    assert sum(len(c) for c in chunks) == 20
    assert all(c for c in chunks)


def test_workspace_has_all_stage_dirs(tmp_path):
    paths = ck.ensure_workspace(str(tmp_path / "course"))
    for name in ck.STAGE_DIRS:
        assert os.path.isdir(paths[name])


def test_collect_pptx_deck_reads_text_tables_notes_and_pictures(tmp_path):
    from pptx import Presentation

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(0, 0, 100000, 100000)
    box.text_frame.text = "Transcranial magnetic stimulation"
    table_shape = slide.shapes.add_table(2, 2, 0, 200000, 400000, 200000)
    table_shape.table.cell(0, 0).text = "TMS"
    table_shape.table.cell(0, 1).text = "磁场刺激"
    buffer = io.BytesIO()
    Image.new("RGB", (300, 200), (10, 20, 30)).save(buffer, format="PNG")
    slide.shapes.add_picture(io.BytesIO(buffer.getvalue()), 0, 400000, 200000, 150000)
    slide.notes_slide.notes_text_frame.text = "记得讲 motor threshold"
    path = tmp_path / "deck.pptx"
    prs.save(str(path))

    deck = ck.collect_pptx_deck(str(path), slides_dir=str(tmp_path / "slides"))
    assert len(deck.slides) == 1
    slide_text = deck.slides[0].plain_text()
    assert "Transcranial magnetic stimulation" in slide_text
    assert "TMS | 磁场刺激" in slide_text
    assert "motor threshold" in slide_text
    assert len(deck.slides[0].images) == 1
    assert os.path.exists(deck.slides[0].images[0])
    assert "<!-- page 1 -->" in deck.markdown()


def test_ledger_lists_files_and_warnings(tmp_path):
    result = ck.CourseJobResult(course_dir=str(tmp_path), slides=3, outline_points=7)
    result.add(str(tmp_path / "a.md"))
    result.warnings.append("模型超时")
    text = ck._ledger(type("C", (), {"course_name": "测试课"})(), "deck.pptx", result)
    assert "测试课" in text
    assert "知识点数量：7" in text
    assert "模型超时" in text


def test_render_markdown_to_docx_roundtrip(tmp_path):
    markdown = (
        "# 标题\n\n## 小节\n\n- 要点一\n- 要点二\n\n"
        "| 考点 | 来源 |\n| --- | --- |\n| Stroop | PPT p.22 |\n\n"
        "这是**粗体**与 `代码` 混排。\n"
    )
    out = tmp_path / "out.docx"
    render_markdown_to_docx(markdown, str(out), title="课程知识总结", subtitle="副标题")
    assert out.exists() and out.stat().st_size > 5000

    from docx import Document

    document = Document(str(out))
    texts = [p.text for p in document.paragraphs]
    assert any("标题" in t for t in texts)
    assert any("要点一" in t for t in texts)
    assert len(document.tables) == 1
    assert document.tables[0].cell(0, 0).text == "考点"
