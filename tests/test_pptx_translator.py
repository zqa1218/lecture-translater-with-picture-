"""幻灯片文本收集与写回：组合图形、表格、备注、首个 run 语义。"""

from __future__ import annotations

from pptx import Presentation
from pptx.util import Emu, Inches

from pptx_translator import apply_translation, collect_text_units, iter_text_units


def _blank():
    prs = Presentation()
    return prs, prs.slides.add_slide(prs.slide_layouts[6])


def _texts(prs):
    return [unit.text for unit in collect_text_units(prs)]


def test_collects_top_level_textbox():
    prs, slide = _blank()
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1)).text_frame.text = "Anode"
    assert _texts(prs) == ["Anode"]


def test_recurses_into_group_shapes():
    prs, slide = _blank()
    group = slide.shapes.add_group_shape()
    group.shapes.add_textbox(Emu(0), Emu(0), Emu(1000), Emu(500)).text_frame.text = "Cathode"
    assert _texts(prs) == ["Cathode"], "组合图形里的文字必须被遍历到"


def test_collects_table_cells():
    prs, slide = _blank()
    table = slide.shapes.add_table(2, 1, Inches(1), Inches(1), Inches(4), Inches(2)).table
    table.cell(0, 0).text_frame.text = "Stimulator"
    table.cell(1, 0).text_frame.text = "Magnetic field"
    assert set(_texts(prs)) == {"Stimulator", "Magnetic field"}


def test_skips_numbers_and_punctuation():
    prs, slide = _blank()
    frame = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1)).text_frame
    frame.text = "2024"
    frame.add_paragraph().text = "  --  "
    frame.add_paragraph().text = "TMS"
    assert _texts(prs) == ["TMS"]


def test_notes_are_left_alone_when_slide_has_none():
    prs, slide = _blank()
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1)).text_frame.text = "Anode"
    assert slide.has_notes_slide is False
    assert _texts(prs) == ["Anode"]
    assert slide.has_notes_slide is False, "不能凭空造出空备注页"


def test_notes_are_collected_when_present():
    prs, slide = _blank()
    slide.notes_slide.notes_text_frame.text = "Remember this for the exam"
    assert _texts(prs) == ["Remember this for the exam"]


def test_apply_translation_keeps_first_run_and_clears_rest():
    _prs, slide = _blank()
    frame = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1)).text_frame
    paragraph = frame.paragraphs[0]
    paragraph.add_run().text = "An"
    paragraph.add_run().text = "ode"
    assert len(paragraph.runs) == 2

    apply_translation(paragraph, "阳极")

    assert paragraph.text == "阳极"
    assert [run.text for run in paragraph.runs] == ["阳极", ""]


def test_apply_translation_adds_a_run_when_there_is_none():
    _prs, slide = _blank()
    frame = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1)).text_frame
    paragraph = frame.paragraphs[0]
    assert list(paragraph.runs) == []

    apply_translation(paragraph, "阴极")

    assert paragraph.text == "阴极"


def test_iter_text_units_yields_paragraph_objects():
    _prs, slide = _blank()
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1)).text_frame.text = "Hello"
    paragraphs = list(iter_text_units(slide.shapes))
    assert len(paragraphs) == 1
    assert paragraphs[0].text == "Hello"