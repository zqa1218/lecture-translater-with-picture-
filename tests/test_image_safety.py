import io
import numpy as np
import pytest
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches

from image_translator import ImagePlan, render_plan
from ocr_engine import OcrEngine, OcrLine, _merge_line
from pipeline import translate_images_in_pptx, translate_pdf
from translator_config import TranslatorConfig


def line(text, box):
    x0, y0, x1, y1 = box
    return OcrLine(((x0,y0),(x1,y0),(x1,y1),(x0,y1)), text, .99)


def plan():
    image = Image.new('RGB', (320, 120), 'white')
    draw = ImageDraw.Draw(image)
    draw.rectangle((20, 30, 80, 50), fill='black')
    draw.rectangle((200, 30, 260, 50), fill='red')
    return ImagePlan(image, np.array(image), False,
                     [line('Anode',(15,25,100,60)), line('MRI',(195,25,280,60))])


def test_only_successful_regions_change_pixels():
    original = plan()
    result, stats = render_plan(original, ['阳极', 'MRI'])
    assert stats.translated == 1
    before, after = original.array, np.array(result)
    mask = np.ones(before.shape[:2], bool)
    mask[20:65, 10:105] = False
    assert np.array_equal(before[mask], after[mask])


def test_tight_ocr_box_does_not_leave_original_text_edges():
    image = Image.new('RGB', (220, 80), 'white')
    draw = ImageDraw.Draw(image)
    # 模拟字形左右各有 2px 落在 OCR 框外。
    draw.rectangle((18, 32, 22, 48), fill='black')
    draw.rectangle((118, 32, 122, 48), fill='black')
    source = ImagePlan(image, np.array(image), False, [line('Original', (20, 30, 120, 50))])
    result, stats = render_plan(source, ['译文'])
    assert stats.translated == 1
    pixels = np.array(result)
    assert np.all(pixels[28:52, 18:20] > 240)
    assert np.all(pixels[28:52, 121:123] > 240)


@pytest.mark.parametrize('translations', [['', 'MRI'], ['太长' * 1000, 'MRI']])
def test_empty_or_oversize_translation_keeps_original(translations):
    result, stats = render_plan(plan(), translations)
    assert result is None and stats.translated == 0


def test_missing_font_does_not_erase(monkeypatch):
    monkeypatch.setattr('image_translator.resolve_font', lambda _: ('',0))
    result, stats = render_plan(plan(), ['阳极', 'MRI'])
    assert result is None and not stats.regions_detail


def test_chinese_native_label_fits_original_box():
    from pdf_builder import _optimize_font_size
    from image_translator import resolve_font
    from PIL import ImageFont
    path,index=resolve_font()
    for text,width,height in [('第二部分',65,20),('T2（x-y 平面内的衰减）',190,24)]:
        size=_optimize_font_size(text,width,height,20)
        font=ImageFont.truetype(path,round(size*4),index=index)
        assert size < 20 and size * 1.25 <= height
        assert font.getlength(text) <= width * 4


def test_subscripts_use_font_with_glyph():
    from image_translator import _font_runs,resolve_font
    from PIL import ImageFont
    path,index=resolve_font()
    font=ImageFont.truetype(path,24,index=index)
    runs=list(_font_runs('T₂',font))
    assert runs[0][1] is font
    assert bytes(runs[1][1].getmask('₂')) != bytes(runs[1][1].getmask(chr(0x10ffff)))


def test_colored_marker_near_black_text_survives_erasure():
    from image_translator import erase_regions
    array=np.full((50,200,3),255,dtype=np.uint8)
    array[10:25,40:100]=0
    array[12:18,12:18]=(255,180,30)
    result=erase_regions(array,[line('Label',(5,5,180,40))],dilate=0)
    assert np.array_equal(result[12:18,12:18],array[12:18,12:18])


def test_thin_colored_text_fringe_is_not_preserved():
    from image_translator import erase_regions
    array=np.full((50,200,3),255,dtype=np.uint8)
    array[15:25,40:100]=0
    array[14,40:100]=(220,40,160)
    array[25,40:100]=(30,120,230)
    result=erase_regions(array,[line('Label',(5,5,180,40))],dilate=0)
    assert np.all(result[14:26,40:100] > 240)


def test_tile_coordinates_and_dedup():
    calls = []
    class Engine:
        def __call__(self, image):
            calls.append(image.shape)
            # 大图首轮空结果；局部块均返回一行。
            return ([[[(20,30),(100,30),(100,50),(20,50)], 'Label', .99]], None) if image.shape[1] <= 1024 else (None,None)
    engine = OcrEngine()
    engine._engine = Engine()
    lines = engine.recognize(Image.new('RGB',(2000,1500)))
    assert any(l.box[0] >= 832 for l in lines)
    assert any(l.box[1] >= 832 for l in lines)
    assert all(0 <= l.box[0] < l.box[2] <= 2000 for l in lines)
    a = [line('Full sentence',(10,10,200,40))]
    _merge_line(a, line('sentence',(90,10,200,40)))
    assert len(a) == 1 and a[0].text == 'Full sentence'


def test_image_geometry_and_unchanged_label_roundtrip(monkeypatch):
    original = plan()
    monkeypatch.setattr('image_translator.prepare_image', lambda *a, **k: original)
    class Client:
        def translate_batch(self, *a, **k):
            return ['阳极','MRI']
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    buf = io.BytesIO(); original.pil.save(buf, format='PNG')
    shape = slide.shapes.add_picture(io.BytesIO(buf.getvalue()), Inches(1), Inches(2), Inches(4))
    shape.crop_left = .1
    shape.rotation = 12
    geometry = (shape.left,shape.top,shape.width,shape.height,shape.crop_left,shape.rotation)
    _, processed, count = translate_images_in_pptx(prs,TranslatorConfig(),Client(),log=lambda _:None)
    assert processed == 1 and count == 1
    out = io.BytesIO(); prs.save(out); out.seek(0)
    shape = Presentation(out).slides[0].shapes[0]
    assert geometry == (shape.left,shape.top,shape.width,shape.height,shape.crop_left,shape.rotation)
    pixels = np.array(Image.open(io.BytesIO(shape.image.blob)))
    assert np.array_equal(pixels[25:60,195:280],original.array[25:60,195:280])


def test_rotated_overlay_keeps_original(monkeypatch):
    original = plan()
    monkeypatch.setattr('image_translator.prepare_image', lambda *a,**k: original)
    class Client:
        def translate_batch(self,*a,**k): return ['阳极','MRI']
    prs=Presentation(); slide=prs.slides.add_slide(prs.slide_layouts[6])
    buf=io.BytesIO(); original.pil.save(buf,format='PNG')
    shape=slide.shapes.add_picture(io.BytesIO(buf.getvalue()),0,0,Inches(4)); shape.rotation=30
    blob=shape.image.blob
    _,processed,_=translate_images_in_pptx(prs,TranslatorConfig(image_mode='overlay'),Client(),log=lambda _:None)
    assert processed == 0 and shape.image.blob == blob and len(slide.shapes)==1


def test_scanned_pdf_reaches_image_translation(tmp_path, monkeypatch):
    import fitz
    doc=fitz.open(); page=doc.new_page(width=320,height=120)
    buf=io.BytesIO(); plan().pil.save(buf,format='PNG')
    page.insert_image(page.rect,stream=buf.getvalue())
    src=tmp_path/'scan.pdf'; doc.save(src); doc.close()
    monkeypatch.setattr('pipeline.make_client',lambda *a,**k: object())
    calls=[]
    def translate(prs,*a,**k):
        calls.append(len(prs.slides)); return 1,0,0
    monkeypatch.setattr('pipeline.translate_images_in_pptx',translate)
    translate_pdf(str(src),str(tmp_path/'out.pptx'),TranslatorConfig(),log=lambda _:None)
    assert calls == [1]
