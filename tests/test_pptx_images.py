"""幻灯片图片几何：组合图形坐标、裁剪、翻转、图片替换。"""

from __future__ import annotations

import io

import pytest
from PIL import Image
from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches

from pptx_images import (
    absolute_box,
    collect_pictures,
    crop_fractions,
    map_pixel_box,
    picture_bytes,
    replace_picture_image,
)


def _png(size=(200, 100), color=(200, 30, 30)):
    buffer = io.BytesIO()
    Image.new("RGB", size, color).save(buffer, format="PNG")
    return buffer.getvalue()


@pytest.fixture()
def deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    picture = slide.shapes.add_picture(io.BytesIO(_png()), Inches(1), Inches(2), Inches(4), Inches(2))
    return prs, slide, picture


def test_absolute_box_of_toplevel_picture(deck):
    _prs, _slide, picture = deck
    left, top, width, height = absolute_box(picture)
    assert (left, top) == (int(Inches(1)), int(Inches(2)))
    assert width == int(Inches(4))
    assert height == int(Inches(2))


def test_map_pixel_box_without_crop_is_linear(deck):
    _prs, _slide, picture = deck
    box = map_pixel_box(picture, (200, 100), (0, 0, 200, 100))
    assert box == (int(Inches(1)), int(Inches(2)), int(Inches(4)), int(Inches(2)))
    half = map_pixel_box(picture, (200, 100), (0, 0, 100, 50))
    assert half[2] == int(Inches(2)) and half[3] == int(Inches(1))


def test_map_pixel_box_respects_crop(deck):
    _prs, _slide, picture = deck
    picture._element.blipFill.append(
        picture._element.blipFill.makeelement(qn("a:srcRect"), {"l": "25000"})
    )
    assert crop_fractions(picture)[0] == pytest.approx(0.25)
    # 可见区域只剩右 3/4；左边缘对应原图 x=50px
    left, _top, width, _height = map_pixel_box(picture, (200, 100), (50, 0, 200, 100))
    assert left == int(Inches(1))
    assert width == int(Inches(4))


def test_map_pixel_box_respects_flip(deck):
    _prs, _slide, picture = deck
    normal = map_pixel_box(picture, (200, 100), (0, 0, 50, 100))
    xfrm = picture._element.spPr.find(qn("a:xfrm"))
    xfrm.set("flipH", "1")
    mirrored = map_pixel_box(picture, (200, 100), (0, 0, 50, 100))
    assert normal != mirrored, "翻转后映射结果必须不同"
    # 水平翻转后，原图左侧 1/4 应落到形状的右侧
    assert mirrored[0] + mirrored[2] == int(Inches(1)) + int(Inches(4))
    assert normal[0] == int(Inches(1))


def test_group_transform_is_applied():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    group = slide.shapes.add_group_shape()
    picture = group.shapes.add_picture(io.BytesIO(_png()), Inches(1), Inches(1), Inches(2), Inches(1))
    # python-pptx 让组合图形的子坐标空间与页面 1:1（chOff/chExt == off/ext）
    left, top, width, height = absolute_box(picture)
    assert (left, top) == (int(Inches(1)), int(Inches(1)))
    assert (width, height) == (int(Inches(2)), int(Inches(1)))
    assert len(collect_pictures(prs)) == 1, "组合图形里的图片也要被遍历到"


def test_group_scaling_is_applied():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    group = slide.shapes.add_group_shape()
    picture = group.shapes.add_picture(io.BytesIO(_png()), Emu(0), Emu(0), Emu(1000), Emu(500))
    xfrm = group._element.find(qn("p:grpSpPr")).find(qn("a:xfrm"))
    xfrm.find(qn("a:ext")).set("cx", "4000")
    xfrm.find(qn("a:ext")).set("cy", "2000")
    # 子坐标 1000 EMU 宽 -> 页面 4000 EMU 宽，放大 4 倍
    _left, _top, width, height = absolute_box(picture)
    assert width == 4000
    assert height == 2000


def test_replace_picture_image_roundtrip(deck):
    prs, _slide, picture = deck
    new_png = _png(size=(64, 64), color=(0, 128, 255))
    assert replace_picture_image(picture, new_png) is True
    buffer = io.BytesIO()
    prs.save(buffer)
    buffer.seek(0)
    reopened = Presentation(buffer)
    blob, extension = picture_bytes(collect_pictures(reopened)[0].shape)
    assert extension == "png"
    assert Image.open(io.BytesIO(blob)).size == (64, 64), "替换后的图片数据必须真的写进了文件"


def test_picture_bytes_handles_missing_image(deck):
    _prs, _slide, picture = deck
    assert picture_bytes(picture)[0] is not None
