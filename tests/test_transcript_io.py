"""课堂素材：时间戳解析、多格式读取、乱序排序、时间轴合并。"""

from __future__ import annotations

import os

import pytest

from transcript_io import (
    Segment,
    format_timestamp,
    merge_segments,
    natural_key,
    order_by_content,
    order_sources,
    parse_timestamp,
    read_material,
    read_text_transcript,
    transcript_duration,
    write_transcript,
)


def test_format_and_parse_timestamp():
    assert format_timestamp(0) == "00:00"
    assert format_timestamp(75) == "01:15"
    assert format_timestamp(3725) == "01:02:05"
    assert parse_timestamp("[12:30] 老师开始讲") == 750
    assert parse_timestamp("(01:02:05)") == 3725
    assert parse_timestamp("没有时间戳") is None


def test_read_plain_transcript_with_timestamps(tmp_path):
    path = tmp_path / "t.txt"
    path.write_text("[00:10] 第一段\n[00:45] 第二段\n\n[01:20] 第三段\n", encoding="utf-8")
    segments = read_text_transcript(str(path))
    assert [s.start for s in segments] == [10, 45, 80]
    assert segments[0].text == "第一段"


def test_read_srt_transcript(tmp_path):
    path = tmp_path / "t.srt"
    path.write_text(
        "1\n00:00:01,000 --> 00:00:04,500\nHello everyone\n\n"
        "2\n00:00:05,000 --> 00:00:08,000\nToday we talk about TMS\n",
        encoding="utf-8",
    )
    segments = read_text_transcript(str(path))
    assert len(segments) == 2
    assert segments[0].start == 1.0 and segments[0].end == 4.5
    assert segments[1].text == "Today we talk about TMS"


def test_read_vtt_transcript(tmp_path):
    path = tmp_path / "t.vtt"
    path.write_text("WEBVTT\n\n00:00:01.000 --> 00:00:03.000\n第一句\n", encoding="utf-8")
    segments = read_text_transcript(str(path))
    assert any(s.text == "第一句" for s in segments)


def test_read_json_transcript(tmp_path):
    path = tmp_path / "t.json"
    path.write_text('{"segments": [{"start": 1, "end": 2, "text": "hello"}]}', encoding="utf-8")
    segments = read_text_transcript(str(path))
    assert segments[0].text == "hello" and segments[0].start == 1


def test_read_docx_transcript(tmp_path):
    from docx import Document

    document = Document()
    document.add_paragraph("[00:05] 老师：今天我们讲 tES")
    document.add_paragraph("第二段内容")
    path = tmp_path / "t.docx"
    document.save(str(path))
    segments = read_text_transcript(str(path))
    assert len(segments) == 2
    assert segments[0].start == 5


def test_natural_key_sorts_numbers_as_numbers():
    names = ["第10段.mp3", "第2段.mp3", "第1段.mp3"]
    assert sorted(names, key=natural_key) == ["第1段.mp3", "第2段.mp3", "第10段.mp3"]


def test_order_sources_by_filename(tmp_path):
    paths = []
    for name in ("part3.wav", "part1.wav", "part2.wav"):
        target = tmp_path / name
        target.write_bytes(b"x")
        paths.append(str(target))
    ordered, method = order_sources(paths)
    assert [os.path.basename(p) for p in ordered] == ["part1.wav", "part2.wav", "part3.wav"]
    assert "文件名" in method


def test_order_sources_respects_explicit_order(tmp_path):
    paths = []
    for name in ("b.wav", "a.wav"):
        target = tmp_path / name
        target.write_bytes(b"x")
        paths.append(str(target))
    ordered, method = order_sources(paths, explicit=True)
    assert ordered == paths
    assert "手动" in method


def test_order_sources_uses_metadata(tmp_path):
    first = tmp_path / "zzz.wav"
    second = tmp_path / "aaa.wav"
    first.write_bytes(b"x")
    second.write_bytes(b"x")
    os.utime(first, (1_600_000_000, 1_600_000_000))
    os.utime(second, (1_700_000_000, 1_700_000_000))
    ordered, _method = order_sources([str(second), str(first)])
    assert ordered[0] == str(first)


class _FakeClient:
    def __init__(self, payload):
        self.payload = payload
        self.prompts = []

    def json_chat(self, prompt, system=None):  # noqa: D102
        self.prompts.append(prompt)
        return self.payload


def test_order_by_content_accepts_valid_permutation():
    client = _FakeClient({"order": [1, 2, 0], "reason": "承接关系", "confidence": "high"})
    order, reason = order_by_content([("a", "开头"), ("b", "中间"), ("c", "结尾")], client)
    assert order == [1, 2, 0]
    assert "承接关系" in reason


def test_order_by_content_keeps_order_when_low_confidence():
    client = _FakeClient({"order": [2, 1, 0], "reason": "不确定", "confidence": "low"})
    order, _reason = order_by_content([("a", "x"), ("b", "y"), ("c", "z")], client)
    assert order == [0, 1, 2]


def test_order_by_content_ignores_invalid_payload():
    client = _FakeClient({"order": [0, 0, 1], "confidence": "high"})
    order, reason = order_by_content([("a", "x"), ("b", "y"), ("c", "z")], client)
    assert order == [0, 1, 2]
    assert "不合法" in reason


def test_merge_segments_offsets_timeline():
    groups = [
        [Segment(text="a1", start=0, end=10), Segment(text="a2", start=10, end=20)],
        [Segment(text="b1", start=0, end=5)],
    ]
    merged = merge_segments(groups)
    assert [s.text for s in merged] == ["a1", "a2", "b1"]
    assert merged[2].start == 20, "第二段的时间要接在第一段后面"
    assert transcript_duration(merged) == 25


def test_write_transcript_roundtrip(tmp_path):
    segments = [Segment(text="第一句", start=5), Segment(text="第二句", start=65)]
    path = write_transcript(str(tmp_path / "out.txt"), segments)
    content = open(path, encoding="utf-8").read()
    assert "[00:05] 第一句" in content
    assert "[01:05] 第二句" in content


def test_read_material_rejects_audio(tmp_path):
    path = tmp_path / "a.mp3"
    path.write_bytes(b"x")
    with pytest.raises(ValueError):
        read_material(str(path))
